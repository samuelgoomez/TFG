"""Genera el Excel con los pares abstract real / abstract generado."""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ORIGINALES_DIR = Path("abstracts/originales")
GENERADOS_DIR = Path("abstracts/generados")
SALIDA = Path("abstracts/comparacion_abstracts.xlsx")


def _extraer_campo(texto: str, etiqueta: str) -> str:
    match = re.search(rf"\*\*{etiqueta}:\*\*\s*(.+)", texto)
    return match.group(1).strip() if match else ""


def _extraer_titulo(texto: str) -> str:
    match = re.search(r"^#\s+(.+)", texto, re.MULTILINE)
    return match.group(1).strip() if match else ""


def _extraer_abstract_original(texto: str) -> str:
    match = re.search(
        r"##\s*(?:Abstract|Resumen) original\s*\n+(.+?)(?:\n##|\Z)",
        texto,
        re.DOTALL,
    )
    return match.group(1).strip() if match else ""


def generar_excel_comparacion() -> Path | None:
    """Recorre abstracts/originales y abstracts/generados y crea el Excel de comparación.

    Devuelve la ruta del Excel generado, o None si no hay ningún par disponible.
    """
    filas = []
    for fichero in sorted(ORIGINALES_DIR.glob("abstract_*.md")):
        nombre_paper = fichero.stem.removeprefix("abstract_")
        generado_path = GENERADOS_DIR / f"{nombre_paper}_generado.md"
        if not generado_path.exists():
            continue

        texto_original = fichero.read_text(encoding="utf-8")
        filas.append({
            "titulo": _extraer_titulo(texto_original),
            "idioma": _extraer_campo(texto_original, "Idioma"),
            "abstract_real": _extraer_abstract_original(texto_original),
            "abstract_generado": generado_path.read_text(encoding="utf-8").strip(),
        })

    if not filas:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparacion abstracts"

    cabeceras = ["Paper", "Idioma", "Abstract real", "Abstract generado"]
    ws.append(cabeceras)
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(vertical="top", wrap_text=True)

    for fila in filas:
        ws.append([fila["titulo"], fila["idioma"], fila["abstract_real"], fila["abstract_generado"]])

    for fila_celdas in ws.iter_rows(min_row=2):
        for celda in fila_celdas:
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    anchos = {"A": 35, "B": 12, "C": 70, "D": 70}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    for i in range(2, len(filas) + 2):
        ws.row_dimensions[i].height = 200

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    return SALIDA
