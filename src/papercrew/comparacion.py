"""Genera los Excel con los pares real / generado para abstracts e introducciones."""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


def _extraer_campo(texto: str, etiqueta: str) -> str:
    match = re.search(rf"\*\*{etiqueta}:\*\*\s*(.+)", texto)
    return match.group(1).strip() if match else ""


def _extraer_titulo(texto: str) -> str:
    match = re.search(r"^#\s+(.+)", texto, re.MULTILINE)
    return match.group(1).strip() if match else ""


def _extraer_seccion_original(texto: str, etiquetas: str) -> str:
    match = re.search(
        rf"##\s*(?:{etiquetas})\s*\n+(.+?)(?:\n##|\Z)",
        texto,
        re.DOTALL,
    )
    return match.group(1).strip() if match else ""


def _generar_excel_pares(
    originales_dir: Path,
    generados_dir: Path,
    salida: Path,
    prefijo_original: str,
    sufijo_generado: str,
    etiquetas_seccion: str,
    nombre_hoja: str,
    cabecera_real: str,
    cabecera_generado: str,
) -> Path | None:
    """Recorre los ficheros originales y generados y crea el Excel de comparación.

    Devuelve la ruta del Excel generado, o None si no hay ningún par disponible.
    """
    filas = []
    for fichero in sorted(originales_dir.glob(f"{prefijo_original}*.md")):
        nombre_paper = fichero.stem.removeprefix(prefijo_original)
        generado_path = generados_dir / f"{nombre_paper}{sufijo_generado}"
        if not generado_path.exists():
            continue

        texto_original = fichero.read_text(encoding="utf-8")
        filas.append({
            "titulo": _extraer_titulo(texto_original),
            "idioma": _extraer_campo(texto_original, "Idioma"),
            "real": _extraer_seccion_original(texto_original, etiquetas_seccion),
            "generado": generado_path.read_text(encoding="utf-8").strip(),
        })

    if not filas:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    cabeceras = ["Paper", "Idioma", cabecera_real, cabecera_generado]
    ws.append(cabeceras)
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(vertical="top", wrap_text=True)

    for fila in filas:
        ws.append([fila["titulo"], fila["idioma"], fila["real"], fila["generado"]])

    for fila_celdas in ws.iter_rows(min_row=2):
        for celda in fila_celdas:
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    anchos = {"A": 35, "B": 12, "C": 70, "D": 70}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    for i in range(2, len(filas) + 2):
        ws.row_dimensions[i].height = 200

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    return salida


def generar_excel_comparacion_abstracts() -> Path | None:
    """Recorre abstracts/originales y abstracts/generados y crea el Excel de comparación de abstracts."""
    return _generar_excel_pares(
        originales_dir=Path("abstracts/originales"),
        generados_dir=Path("abstracts/generados"),
        salida=Path("abstracts/comparacion_abstracts.xlsx"),
        prefijo_original="abstract_",
        sufijo_generado="_generado.md",
        etiquetas_seccion="Abstract original|Resumen original",
        nombre_hoja="Comparacion abstracts",
        cabecera_real="Abstract real",
        cabecera_generado="Abstract generado",
    )


def generar_excel_comparacion_introducciones() -> Path | None:
    """Recorre introducciones/originales y introducciones/generadas y crea el Excel de comparación de introducciones."""
    return _generar_excel_pares(
        originales_dir=Path("introducciones/originales"),
        generados_dir=Path("introducciones/generadas"),
        salida=Path("introducciones/comparacion_introducciones.xlsx"),
        prefijo_original="introduccion_",
        sufijo_generado="_generada.md",
        etiquetas_seccion="Introducción original|Introduccion original",
        nombre_hoja="Comparacion introducciones",
        cabecera_real="Introducción real",
        cabecera_generado="Introducción generada",
    )
